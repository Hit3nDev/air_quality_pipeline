"""
Builds AirLens_Project_Report.docx from the live database and report charts.
Run:  python make_report.py   (after at least one pipeline run)
"""
import sqlite3
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

BASE = Path(__file__).parent
DB, REPORT_DIR = BASE / "data" / "air_quality.db", BASE / "reports"
OUT = BASE / "AirLens_Project_Report.docx"
LIMIT = 60

TEAM = [("Hiten Mandhyan", "24215117"), ("Naman Jiwnani", "24215125"), ("Tejas Kashyap", "24215043")]
CLASS, SUBJECT, UNIV = "5 BCA A", "Data Engineering", "Christ University, Delhi NCR"

# ---------- live figures ----------
def load():
    if not DB.exists():
        return None
    with sqlite3.connect(DB) as con:
        n = pd.read_sql("""SELECT (SELECT COUNT(*) FROM raw_readings) raw, (SELECT COUNT(*) FROM clean_readings) clean,
                           (SELECT COUNT(*) FROM clean_readings WHERE is_outlier=1) outliers,
                           (SELECT COUNT(*) FROM clean_readings WHERE was_imputed=1) imputed,
                           (SELECT COUNT(*) FROM daily_summary) daily, (SELECT COUNT(*) FROM archive_readings) archive,
                           (SELECT COUNT(*) FROM pipeline_runs) runs""", con).iloc[0]
        daily = pd.read_sql("SELECT * FROM daily_summary", con, parse_dates=["date"])
        hourly = pd.read_sql("SELECT city,time,pm2_5 FROM clean_readings WHERE is_outlier=0", con, parse_dates=["time"])
    if daily.empty:
        return None
    stats = daily.groupby("city").agg(mean=("avg_pm2_5", "mean"), worst=("avg_pm2_5", "max"))
    stats["over"] = daily.assign(o=daily.avg_pm2_5 > LIMIT).groupby("city").o.mean() * 100
    hourly["hour"] = hourly.time.dt.hour
    stats["peak_hour"] = hourly.groupby(["city", "hour"]).pm2_5.mean().unstack().idxmax(axis=1)
    stats = stats.sort_values("mean", ascending=False).round(1)
    return dict(n=n, daily=daily, stats=stats, start=daily.date.min(), end=daily.date.max())

# ---------- document helpers ----------
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5); s.left_margin = s.right_margin = Cm(2.5)
base = doc.styles["Normal"]; base.font.name = "Calibri"; base.font.size = Pt(11)
base.paragraph_format.space_after = Pt(6); base.paragraph_format.line_spacing = 1.15
for lvl, size in [(1, 16), (2, 13)]:
    st = doc.styles[f"Heading {lvl}"]; st.font.name = "Calibri"; st.font.size = Pt(size)
    st.font.bold = True; st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

def para(text, bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullets(items):
    for it in items: doc.add_paragraph(it, style="List Bullet")

def table(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph()

def figure(path, caption, width=15):
    if Path(path).exists():
        doc.add_picture(str(path), width=Cm(width)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_page_number():
    p = doc.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); f1, f2, f3 = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin"); f2.text = "PAGE"; f3.set(qn("w:fldCharType"), "end")
    for e in (f1, f2, f3): r._r.append(e)

# ---------- cover page ----------
d = load()
for _ in range(6): doc.add_paragraph()
para(UNIV, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(f"{SUBJECT} — Project Report", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
para("AirLens", bold=True, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("An Air Quality Monitoring Pipeline Implementing the Complete Data Lifecycle",
     size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
para("Submitted by", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
for name, reg in TEAM: para(f"{name}  ({reg})", align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(CLASS, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
para(datetime.now().strftime("%d %B %Y"), align=WD_ALIGN_PARAGRAPH.CENTER)
add_page_number()
page_break()

# ---------- 1 abstract ----------
doc.add_heading("1. Abstract", 1)
para("AirLens is an end-to-end data engineering project that monitors air pollution in five Indian cities: Delhi, Mumbai, "
     "Bengaluru, Kolkata and Chennai. It implements every stage of the data lifecycle on a live dataset. Hourly PM2.5, PM10, "
     "NO₂ and AQI readings are collected from a public API, stored in a relational database, cleaned, aggregated, analysed "
     "against India's legal pollution limit, published through an interactive dashboard, and archived under a defined "
     "retention policy. The project demonstrates that the data lifecycle is not an abstract model but a sequence of concrete, "
     "automatable and auditable engineering steps.")

# ---------- 2 introduction ----------
doc.add_heading("2. Introduction and Objectives", 1)
para("Air pollution is a persistent public-health problem in Indian cities, with Delhi regularly recording PM2.5 levels "
     "several times the national standard. Data on this is freely available but arrives hourly, with gaps and sensor errors, "
     "and in a form that is not directly useful to a reader. Turning such raw data into reliable, understandable information "
     "is the central task of data engineering. The objectives of this project were to:")
bullets(["Build an automated pipeline that implements collection, storage, cleaning, transformation, analysis, sharing and archival on real data.",
         "Make each stage observable, so the pipeline records what it did and can be audited.",
         "Present the results in a dashboard that communicates findings clearly to a non-technical reader.",
         "Apply governance practices such as retention policies, run logging and data-quality measurement."])

# ---------- 3 dataset ----------
doc.add_heading("3. Data Source and Dataset", 1)
para("Data is obtained from the Open-Meteo Air Quality API, which provides hourly pollutant concentrations derived from the "
     "Copernicus Atmosphere Monitoring Service (CAMS) and requires no authentication. The API allows up to 92 days of past "
     "data in one request, which the pipeline uses for its initial backfill. The reference threshold used throughout is "
     f"India's National Ambient Air Quality Standard for PM2.5 of {LIMIT} µg/m³ averaged over 24 hours.")
table(["Variable", "Unit", "Description"],
      [["pm2_5", "µg/m³", "Fine particulate matter below 2.5 µm"],
       ["pm10", "µg/m³", "Particulate matter below 10 µm"],
       ["nitrogen_dioxide", "µg/m³", "NO₂, mainly from vehicle exhaust"],
       ["us_aqi", "index", "US Air Quality Index, 0 to 500"]], widths=[4, 2.5, 9])
if d:
    para(f"The dataset currently covers {d['start']:%d %B %Y} to {d['end']:%d %B %Y}, with {int(d['n'].raw):,} raw hourly "
         f"records in the working store, {int(d['n'].archive):,} archived records and {int(d['n'].daily):,} city-day aggregates.")

# ---------- 4 architecture ----------
doc.add_heading("4. System Architecture and Tools", 1)
para("The system has two components: a batch pipeline (pipeline.py) that runs on a schedule, and a dashboard (dashboard.py) "
     "that reads from the same database on demand. SQLite was chosen because it is file-based, needs no server, and still "
     "provides SQL, primary keys and upserts; the design is portable to PostgreSQL without changes to the pipeline logic.")
table(["Layer", "Tool", "Role"],
      [["Collection", "Python requests", "Calls the Open-Meteo API per city"],
       ["Raw storage", "JSON files + SQLite", "Immutable API snapshots and raw_readings table"],
       ["Processing", "pandas", "Cleaning, imputation, outlier detection, aggregation"],
       ["Analytics storage", "SQLite", "clean_readings, daily_summary, archive_readings"],
       ["Reporting", "matplotlib, Markdown", "Written summary and charts in reports/"],
       ["Sharing", "Streamlit, Plotly", "Interactive web dashboard"],
       ["Governance", "pipeline_runs table, logging", "Audit log of every run"],
       ["Scheduling", "cron / Task Scheduler", "Daily execution"]], widths=[3.5, 4.5, 7.5])

# ---------- 5 lifecycle ----------
doc.add_heading("5. Implementation of the Data Lifecycle", 1)
stages = [
    ("Stage 1: Collect", "For each city the pipeline requests hourly readings with past_days=92 on the first run and "
     "past_days=2 on later runs. Re-fetching the last two days is intentional: it picks up values revised by the provider."),
    ("Stage 2: Store", "Every API response is written unchanged to data/raw/ as a timestamped JSON file so the original data "
     "can always be reprocessed. Parsed rows are upserted into raw_readings on a composite key of city and timestamp, which "
     "makes the pipeline idempotent: running it twice never duplicates data."),
    ("Stage 3: Clean", "Forecast hours later than the current time are dropped; negative concentrations, which are physically "
     "impossible, are set to null; gaps of three hours or less are filled by linear interpolation and flagged was_imputed; "
     "longer gaps are dropped; and any PM2.5 value more than three standard deviations from the city's centred seven-day "
     "rolling mean is flagged is_outlier. Outliers are kept in the store for transparency but excluded from analysis."),
    ("Stage 4: Transform", "Valid hourly rows are aggregated into daily_summary, one row per city per day, with mean and "
     "maximum PM2.5, mean PM10 and NO₂, mean AQI, contributing hours and a categorical AQI label. This table is upserted, "
     "never truncated, so it accumulates history indefinitely."),
    ("Stage 5: Analyse", "A Markdown report and two charts are regenerated on every run, covering city rankings, the worst "
     "city-days, the share of days above the limit, weekday versus weekend averages and the peak hour for each city."),
    ("Stage 6: Share", "The Streamlit dashboard reads directly from the database and offers CSV and Markdown downloads, so "
     "results are available interactively and as portable files."),
    ("Stage 7: Archive and govern", "Hourly rows older than 60 days move from raw_readings to archive_readings; raw JSON "
     "older than 30 days is deleted; daily aggregates are kept permanently; the database is compacted with VACUUM. Every "
     "execution writes a record to pipeline_runs with timing, row counts and status, alongside a text log."),
]
for title, body in stages:
    doc.add_heading(title, 2); para(body)
if d:
    n = d["n"]
    table(["Stage", "Live figure from the database"],
          [["Raw rows stored", f"{int(n.raw):,}"], ["Rows after cleaning", f"{int(n.clean):,}"],
           ["Short gaps filled", f"{int(n.imputed):,}"], ["Outliers flagged", f"{int(n.outliers):,}"],
           ["Daily aggregates", f"{int(n.daily):,}"], ["Rows archived", f"{int(n.archive):,}"],
           ["Pipeline runs logged", f"{int(n.runs):,}"]], widths=[7, 8.5])

# ---------- 6 dashboard ----------
doc.add_heading("6. Dashboard and Visualisation", 1)
para("The dashboard is designed as an air-quality bulletin rather than a generic admin panel. It opens with a headline "
     "stating the worst current reading relative to the legal limit and a plain-language health advisory. Below this, "
     "haze bands show the last seven days for each city hour by hour, coloured on a single scale from blue (clear air) "
     "through amber (exactly the legal limit) to brown and soot. The same scale is reused in every chart.")
para("The lifecycle is visualised in two ways: a rail of seven stages showing live figures from the database and file "
     "system, and a row ledger showing where every collected row ends up, with losses at the cleaning and outlier stages "
     "highlighted so the arithmetic from collection to analysis is visible. Further sections provide per-city trend panels "
     "against the limit, an hour-by-weekday heatmap, a one-square-per-day calendar, a data-quality table, the run audit "
     "log and download buttons.")
para("Insert dashboard screenshots here (headline and haze bands; lifecycle rail and ledger; trend panels).", italic=True, size=9)

# ---------- 7 results ----------
doc.add_heading("7. Results and Key Findings", 1)
if d:
    s = d["stats"]
    para(f"Figures below are computed from the database for {d['start']:%d %b %Y} to {d['end']:%d %b %Y}.")
    table(["City", "Mean PM2.5 (µg/m³)", "Worst day (µg/m³)", "Days above limit", "Peak hour"],
          [[c, r["mean"], r["worst"], f"{r['over']:.0f} %", f"{int(r['peak_hour']):02d}:00"] for c, r in s.iterrows()],
          widths=[3, 3.5, 3.2, 3, 2.5])
    top, low = s.index[0], s.index[-1]
    bullets([f"{top} recorded the highest mean PM2.5 ({s.loc[top,'mean']} µg/m³) and exceeded the legal limit on "
             f"{s.loc[top,'over']:.0f} % of days, with a worst daily average of {s.loc[top,'worst']} µg/m³.",
             f"{low} was the cleanest city, averaging {s.loc[low,'mean']} µg/m³ and exceeding the limit on only "
             f"{s.loc[low,'over']:.0f} % of days.",
             f"Pollution in {top} peaks around {int(s.loc[top,'peak_hour']):02d}:00, a pattern consistent with night-time "
             "temperature inversions that trap emissions near the ground.",
             f"Data quality from the source was high: only {int(d['n'].imputed)} hourly values needed interpolation and "
             f"{int(d['n'].outliers)} were flagged as statistical outliers out of {int(d['n'].clean):,} cleaned rows."])
else:
    para("Run the pipeline before generating the report to populate this section with live figures.", italic=True)
figure(REPORT_DIR / "daily_pm25.png", "Figure 1. Daily average PM2.5 per city against the NAAQS limit.")
figure(REPORT_DIR / "hourly_profile.png", "Figure 2. Average PM2.5 by hour of day.", width=13)

# ---------- 8 challenges ----------
doc.add_heading("8. Challenges, Limitations and Future Work", 1)
doc.add_heading("Challenges", 2)
bullets(["The API returns forecast hours mixed with observations, which had to be filtered out at the cleaning stage.",
         "The retention policy had to be designed so that archiving hourly rows did not destroy the daily aggregates derived from them.",
         "The dashboard had to handle empty or partial states gracefully, such as before the first backfill or while the archive is still empty."])
doc.add_heading("Limitations", 2)
bullets(["Readings are model-derived from CAMS rather than ground sensors and may differ from CPCB station data.",
         "Only five cities and four pollutants are covered; the 92-day API limit means earlier history cannot be recovered.",
         "SQLite is adequate for this volume but not for concurrent writers."])
doc.add_heading("Future work", 2)
bullets(["Integrate CPCB ground-station data for validation.",
         "Add threshold alerts by e-mail or messaging when a city crosses the limit.",
         "Migrate storage to PostgreSQL and scheduling to Apache Airflow.",
         "Add a festival calendar to quantify Diwali-period spikes, and extend coverage to more cities and pollutants."])

# ---------- 9 conclusion + references ----------
doc.add_heading("9. Conclusion", 1)
para("The project shows that the full data lifecycle can be implemented on a small, live dataset with open tools, and "
     "that doing so turns noisy hourly feeds into a trustworthy, explainable picture of air quality. The most valuable "
     "lessons were engineering ones: making the pipeline idempotent, keeping raw data immutable, separating detail from "
     "aggregates so retention does not destroy history, and recording every run so the system can describe itself. "
     "These practices scale directly to larger production data systems.")
doc.add_heading("References", 1)
for ref in ["Open-Meteo, Air Quality API documentation, https://open-meteo.com/en/docs/air-quality-api",
            "Copernicus Atmosphere Monitoring Service (CAMS), ECMWF",
            "Central Pollution Control Board, National Ambient Air Quality Standards, 2009",
            "United States EPA, Technical Assistance Document for the Reporting of Daily Air Quality",
            "pandas, Streamlit and Plotly official documentation"]:
    doc.add_paragraph(ref, style="List Number")

doc.save(OUT)
print(f"Report written to {OUT}")